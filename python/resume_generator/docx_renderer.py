"""ATS-safe plain DOCX resume export (no tables/columns)."""

from __future__ import annotations

from pathlib import Path


def render_resume_docx(ctx: dict, output_path: Path) -> bool:
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        return False

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    name = ctx.get("name") or "Candidate"
    doc.add_heading(name, level=0)
    if ctx.get("subtitle"):
        doc.add_paragraph(ctx["subtitle"])
    if ctx.get("contact"):
        doc.add_paragraph(ctx["contact"])

    summary = ctx.get("summary") or ctx.get("about_me") or ""
    if summary:
        doc.add_heading("Professional Summary", level=1)
        doc.add_paragraph(summary)

    experience = ctx.get("experience") or []
    if experience:
        doc.add_heading("Experience", level=1)
        for job in experience:
            header = f"{job.get('company', '')} — {job.get('title', '')}"
            dates = job.get("dates") or ""
            p = doc.add_paragraph()
            p.add_run(header).bold = True
            if dates:
                p.add_run(f" ({dates})")
            for bullet in job.get("bullets") or []:
                doc.add_paragraph(bullet, style="List Bullet")

    skills = ctx.get("skill_categories") or []
    if skills:
        doc.add_heading("Skills", level=1)
        for cat in skills:
            label = cat.get("label") or "Skills"
            items = ", ".join(cat.get("skills") or [])
            if items:
                doc.add_paragraph(f"{label}: {items}")

    education = ctx.get("education")
    if education:
        doc.add_heading("Education", level=1)
        if isinstance(education, dict):
            doc.add_paragraph(education.get("degree") or education.get("school") or str(education))
        else:
            doc.add_paragraph(str(education))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return True
